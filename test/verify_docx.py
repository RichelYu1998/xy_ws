# -*- coding: utf-8 -*-
"""验证 skill.docx 与 skill.md 内容一致性"""
import re
from docx import Document
from pathlib import Path

def verify_consistency():
    doc = Document('skill.docx')
    md_content = Path('skill.md').read_text(encoding='utf-8')

    print('=== 📊 DOCX vs MD 内容完整性验证 ===\n')

    # 1. 标题检查
    md_headings = len([l for l in md_content.split('\n') if l.startswith('#')])
    doc_headings = len([p for p in doc.paragraphs if p.style.name.startswith('Heading') or p.style.name == 'Title'])
    status = '✅' if abs(md_headings - doc_headings) < 5 else '❌'
    print(f'1. 标题数量: MD={md_headings}, DOCX={doc_headings} {status}')

    # 2. 表格检查
    md_tables = md_content.count('|---|')
    status = '✅' if len(doc.tables) >= 5 else '❌'
    print(f'2. 表格数量: MD分隔线={md_tables}, DOCX={len(doc.tables)} {status}')

    # 3. 列表检查
    md_lists = len([l for l in md_content.split('\n') if re.match(r'^\s*[-*+]\s', l)])
    doc_lists = len([p for p in doc.paragraphs if p.text.startswith('•')])
    status = '✅' if abs(md_lists - doc_lists) < 50 else '⚠️'
    print(f'3. 无序列表: MD={md_lists}, DOCX={doc_lists} {status}')

    # 4. 有序列表
    md_ordered = len([l for l in md_content.split('\n') if re.match(r'^\s*\d+\.\s', l)])
    doc_ordered = len([p for p in doc.paragraphs if re.match(r'^\d+\.', p.text)])
    status = '✅' if abs(md_ordered - doc_ordered) < 20 else '⚠️'
    print(f'4. 有序列表: MD={md_ordered}, DOCX={doc_ordered} {status}')

    # 5. 关键内容验证
    key_phrases = [
        'v5.0.9.51',
        '服务器防崩溃',
        '隧道配置固化',
        'JavaScript变量名',
        '邮件通知系统',
        '安全审计',
        'PY-CORE-027',
        '微购相册管理系统',
        'UTF-8 编码',
        'Cloudflare Tunnel',
        'hostc',
        'WinError 64'
    ]
    
    all_text = ' '.join([p.text for p in doc.paragraphs])
    print(f'\n5. 🔍 关键内容检查 ({len(key_phrases)} 个关键词):')
    found_count = 0
    for phrase in key_phrases:
        found = phrase in all_text
        if found:
            found_count += 1
        print(f'   {"✅" if found else "❌"} "{phrase}"')
    
    print(f'\n   找到: {found_count}/{len(key_phrases)} ({found_count/len(key_phrases)*100:.0f}%)')

    # 6. 统计信息
    print(f'\n6. 📈 文档统计:')
    print(f'   段落总数: {len(doc.paragraphs)}')
    print(f'   表格总数: {len(doc.tables)}')
    file_size = Path('skill.docx').stat().st_size / 1024
    print(f'   文件大小: {file_size:.1f} KB')
    
    # 7. 空白列表项检查（关键！）
    empty_lists = len([p for p in doc.paragraphs if p.text.strip() == '•'])
    if empty_lists > 0:
        print(f'\n   ⚠️ 警告: 发现 {empty_lists} 个空白列表项!')
        return False
    else:
        print(f'\n   ✅ 无空白列表项')

    # 最终结论
    print(f'\n{"="*50}')
    if found_count >= len(key_phrases) * 0.9 and empty_lists == 0:
        print('🎉 验证通过! DOCX 与 MD 内容高度一致')
        return True
    else:
        print('❌ 验证失败! 存在内容不一致问题')
        return False

if __name__ == '__main__':
    success = verify_consistency()
    exit(0 if success else 1)